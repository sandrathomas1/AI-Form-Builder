"""Excel sheet -> draft spec: one printed page per visible worksheet."""

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ai_form_builder.services.extract.fields import mark_fields

MAX_ROWS = 300
MAX_COLS = 40


def _text(value) -> str:
	if value is None:
		return ""
	if hasattr(value, "strftime"):
		return value.strftime("%d-%m-%Y")
	if isinstance(value, float) and value.is_integer():
		return str(int(value))
	return str(value).strip()


def _bg(cell) -> str | None:
	fill = cell.fill
	if not fill or fill.fill_type != "solid":
		return None
	rgb = getattr(fill.fgColor, "rgb", None)
	if not isinstance(rgb, str) or len(rgb) not in (6, 8):
		return None
	rgb = rgb[-6:].upper()
	return None if rgb in ("FFFFFF", "000000") else f"#{rgb}"


def _used_range(sheet):
	last_row = last_col = 0
	for row in sheet.iter_rows(max_row=min(sheet.max_row, MAX_ROWS), max_col=min(sheet.max_column, MAX_COLS)):
		for cell in row:
			if cell.value not in (None, "") or cell.border.left.style or cell.border.bottom.style:
				last_row = max(last_row, cell.row)
				last_col = max(last_col, cell.column)
	for merged in sheet.merged_cells.ranges:
		if merged.min_row <= MAX_ROWS and merged.min_col <= MAX_COLS:
			last_row = max(last_row, min(merged.max_row, MAX_ROWS))
			last_col = max(last_col, min(merged.max_col, MAX_COLS))
	return last_row, last_col


def read_sheet(sheet) -> dict | None:
	last_row, last_col = _used_range(sheet)
	if not last_row or not last_col:
		return None
	anchors, covered = {}, set()
	for merged in sheet.merged_cells.ranges:
		if merged.min_row > last_row or merged.min_col > last_col:
			continue
		anchors[(merged.min_row, merged.min_col)] = (
			min(merged.max_row, last_row) - merged.min_row + 1,
			min(merged.max_col, last_col) - merged.min_col + 1,
		)
		for r in range(merged.min_row, min(merged.max_row, last_row) + 1):
			for c in range(merged.min_col, min(merged.max_col, last_col) + 1):
				if (r, c) != (merged.min_row, merged.min_col):
					covered.add((r, c))
	widths = []
	for col in range(1, last_col + 1):
		dimension = sheet.column_dimensions.get(get_column_letter(col))
		width = dimension.width if dimension is not None and dimension.width else 8.43
		widths.append(0 if dimension is not None and dimension.hidden else round(width, 2))
	rows = []
	for r in range(1, last_row + 1):
		if sheet.row_dimensions.get(r) is not None and sheet.row_dimensions[r].hidden:
			continue
		row = []
		for c in range(1, last_col + 1):
			if (r, c) in covered or not widths[c - 1]:
				continue
			cell = sheet.cell(row=r, column=c)
			item = {"text": _text(cell.value)}
			rowspan, colspan = anchors.get((r, c), (1, 1))
			if colspan > 1:
				item["colspan"] = colspan
			if rowspan > 1:
				item["rowspan"] = rowspan
			if cell.font is not None and cell.font.b:
				item["bold"] = True
			if cell.alignment is not None and cell.alignment.horizontal in ("center", "right", "left"):
				item["align"] = cell.alignment.horizontal
			bg = _bg(cell)
			if bg:
				item["bg"] = bg
			row.append(item)
		if row:
			rows.append(row)
	return {"type": "table", "widths": [w for w in widths if w], "rows": rows}


def _first_bold(block: dict) -> str:
	"""The sheet's heading: its first bold caption, as printed."""
	for row in block["rows"][:6]:
		for cell in row:
			if cell.get("bold") and cell.get("text"):
				return cell["text"]
	return ""


def read_workbook(path: str) -> dict:
	workbook = load_workbook(path, data_only=True)
	pages, title, landscape = [], "", False
	for sheet in workbook.worksheets:
		if sheet.sheet_state != "visible":
			continue
		block = read_sheet(sheet)
		if not block:
			continue
		title = title or _first_bold(block) or sheet.title
		pages.append({"blocks": [block]})
		landscape = landscape or sheet.page_setup.orientation == "landscape"
	spec = {
		"title": title,
		"page": {"size": "A4", "orientation": "landscape" if landscape else "portrait"},
		"fields": [],
		"grids": [],
		"pages": pages,
	}
	return mark_fields(spec)
