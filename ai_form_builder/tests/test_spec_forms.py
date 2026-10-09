"""Spec forms: a client sheet becomes a form that is data, never a DocType.

Run with ``bench --site <site> run-tests --app ai_form_builder``. The AI is
mocked; the Excel and Word readers run for real.
"""

import io
import json
import zipfile
from unittest.mock import patch

import fitz
import frappe
from frappe.tests.utils import FrappeTestCase
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ai_form_builder.api import project_configuration, records, spec_forms
from ai_form_builder.api import template as template_api
from ai_form_builder.services import form_registry, form_spec, spec_renderer, spec_service
from ai_form_builder.services.extract import docx as docx_reader
from ai_form_builder.services.extract import xlsx as xlsx_reader
from ai_form_builder.tests.test_form_library import PROJECT, FormLibraryTestCase, _pdf_url

NO_AI = {"provider": None, "api_key": None, "model": None, "source": None}


def _workbook_bytes():
	book = Workbook()
	sheet = book.active
	sheet.title = "Pour Checklist"
	sheet["A1"] = "CONCRETE POUR CHECKLIST"
	sheet["A1"].font = Font(bold=True)
	sheet.merge_cells("A1:D1")
	sheet["A2"] = "Project No:"
	sheet["C2"] = "Date"
	sheet["A3"] = "Inspected by ________"
	sheet["C3"] = "Remarks"
	sheet["A4"] = "Item"
	sheet["A4"].fill = PatternFill("solid", fgColor="FFD9D9D9")
	sheet["B4"] = "Result"
	sheet["B4"].fill = PatternFill("solid", fgColor="FFD9D9D9")
	sheet.column_dimensions["A"].width = 30
	stream = io.BytesIO()
	book.save(stream)
	return stream.getvalue()


def _file_url(name, content):
	return (
		frappe.get_doc({"doctype": "File", "file_name": name, "content": content, "is_private": 1})
		.insert(ignore_permissions=True)
		.file_url
	)


def _docx_bytes():
	w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
	document = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document {w}><w:body>
<w:p><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:rPr><w:b/></w:rPr><w:t>PERMIT TO WORK</w:t></w:r></w:p>
<w:tbl><w:tblGrid><w:gridCol w:w="3000"/><w:gridCol w:w="6000"/></w:tblGrid>
<w:tr><w:tc><w:tcPr><w:gridSpan w:val="2"/><w:shd w:fill="D9D9D9"/></w:tcPr><w:p><w:r><w:t>SECTION A</w:t></w:r></w:p></w:tc></w:tr>
<w:tr><w:tc><w:p><w:r><w:t>Permit No</w:t></w:r></w:p></w:tc><w:tc><w:p/></w:tc></w:tr>
<w:tr><w:tc><w:p><w:r><w:t>Issuer Signature</w:t></w:r></w:p></w:tc><w:tc><w:p/></w:tc></w:tr>
</w:tbl>
<w:sectPr><w:pgSz w:orient="landscape"/></w:sectPr>
</w:body></w:document>"""
	stream = io.BytesIO()
	with zipfile.ZipFile(stream, "w") as archive:
		archive.writestr("word/document.xml", document)
	return stream.getvalue()


def _custom_doctype_count():
	return frappe.db.count("DocType", {"custom": 1}), frappe.db.count("Client Script")


class TestSpecFormat(FrappeTestCase):
	def test_validate_lists_every_problem(self):
		spec = form_spec.normalise(
			{
				"title": "X",
				"fields": [{"name": "a", "label": "A"}, {"name": "a", "label": "A again"}, {"name": "1bad"}],
				"grids": [{"name": "rows", "columns": []}],
				"pages": [
					{
						"blocks": [
							{"type": "table", "rows": [[{"field": "missing"}]]},
							{"type": "grid", "grid": "nope"},
						]
					}
				],
			}
		)
		errors = " | ".join(form_spec.validate(spec))
		for expected in ("used twice", "'1bad'", "no columns", "'missing'", "'nope'"):
			self.assertIn(expected, errors)
		with self.assertRaises(form_spec.SpecError):
			form_spec.clean(spec)

	def test_model_answer_becomes_the_stored_shape(self):
		payload = {
			"title": "T",
			"fields": [
				{
					"name": "who",
					"label": "Who",
					"type": "Data",
					"options": [],
					"mandatory": False,
					"sample": "",
					"section": "",
				}
			],
			"grids": [
				{
					"name": "items",
					"label": "Items",
					"section": "",
					"columns": [
						{"name": "item", "label": "Item", "type": "Data", "width": 70, "options": []}
					],
					"fixed_rows": [[{"column": "item", "text": "Formwork clean"}]],
					"blank_rows": 3,
					"allow_add": True,
				}
			],
			"pages": [
				{
					"blocks": [
						{
							"type": "table",
							"widths": None,
							"rows": [
								[
									{
										"text": "Who",
										"field": "who",
										"colspan": None,
										"rowspan": None,
										"bold": None,
										"align": None,
										"bg": None,
									}
								]
							],
							"grid": None,
							"text": None,
							"bold": None,
							"align": None,
							"size_pt": None,
							"height_mm": None,
						},
						{
							"type": "grid",
							"widths": None,
							"rows": None,
							"grid": "items",
							"text": None,
							"bold": None,
							"align": None,
							"size_pt": None,
							"height_mm": None,
						},
					]
				}
			],
		}
		spec = form_spec.from_llm(payload)
		self.assertEqual(form_spec.validate(spec), [])
		self.assertEqual(spec["grids"][0]["fixed_rows"], [{"item": "Formwork clean"}])
		self.assertEqual(spec["pages"][0]["blocks"][0]["rows"][0][0], {"text": "Who", "field": "who"})

	def test_schema_is_accepted_by_structured_outputs(self):
		"""Every object closed and fully required; no recursion."""

		def walk(node):
			if isinstance(node, dict):
				if node.get("type") == "object":
					self.assertIs(node.get("additionalProperties"), False)
					self.assertEqual(sorted(node["required"]), sorted(node["properties"]))
				for value in node.values():
					walk(value)
			elif isinstance(node, list):
				for value in node:
					walk(value)

		walk(form_spec.LLM_SCHEMA)
		json.dumps(form_spec.LLM_SCHEMA)


class TestSheetReaders(FrappeTestCase):
	def test_excel_cells_merges_and_fields(self):
		path = frappe.get_site_path("private", "files", "afb-test-reader.xlsx")
		with open(path, "wb") as handle:
			handle.write(_workbook_bytes())
		spec = xlsx_reader.read_workbook(path)
		self.assertEqual(spec["title"], "CONCRETE POUR CHECKLIST")
		rows = spec["pages"][0]["blocks"][0]["rows"]
		self.assertEqual(rows[0][0], {"text": "CONCRETE POUR CHECKLIST", "colspan": 4, "bold": True})
		fields = {f["name"]: f for f in spec["fields"]}
		self.assertEqual(fields["project_no"]["type"], "Data")
		self.assertEqual(fields["date"]["type"], "Date")
		self.assertEqual(fields["inspected_by"]["label"], "Inspected by")
		self.assertEqual(fields["remarks"]["type"], "Text")
		# A shaded header box is a caption, never a field.
		self.assertEqual(rows[3][1].get("bg"), "#D9D9D9")
		self.assertNotIn("field", rows[3][1])
		self.assertEqual(form_spec.validate(form_spec.normalise(spec)), [])

	def test_word_tables_spans_and_fields(self):
		path = frappe.get_site_path("private", "files", "afb-test-reader.docx")
		with open(path, "wb") as handle:
			handle.write(_docx_bytes())
		spec = docx_reader.read_document(path)
		self.assertEqual(spec["title"], "PERMIT TO WORK")
		self.assertEqual(spec["page"]["orientation"], "landscape")
		table = spec["pages"][0]["blocks"][1]
		self.assertEqual(table["widths"], [3000, 6000])
		self.assertEqual(table["rows"][0][0]["colspan"], 2)
		self.assertEqual(table["rows"][0][0]["bg"], "#D9D9D9")
		self.assertEqual(
			{f["name"]: f["type"] for f in spec["fields"]},
			{"permit_no": "Data", "issuer_signature": "Signature"},
		)


class TestSpecForms(FormLibraryTestCase):
	def setUp(self):
		super().setUp()
		# Forms link to the test project, not Docuflow's DCMS Project.
		frappe.db.set_single_value("AI Form Builder Settings", "default_reference_doctype", PROJECT)

	def tearDown(self):
		frappe.set_user("Administrator")
		frappe.db.set_single_value("AI Form Builder Settings", "default_reference_doctype", None)
		for name in frappe.get_all(
			"AI Form Record", filters={"form_code": ["like", "AFB-T-%"]}, pluck="name"
		):
			frappe.delete_doc("AI Form Record", name, force=True, ignore_permissions=True)
		super().tearDown()

	def _excel_form(self, code="AFB-T-SPEC"):
		created = template_api.create_template_from_upload(
			"AFB Test Pour Checklist",
			code,
			_file_url(f"{code}.xlsx", _workbook_bytes()),
			target_area="Quality",
			form_group="QC",
			enable_project_configuration=1,
		)
		with patch("ai_form_builder.integrations.llm.get_config", return_value=NO_AI):
			spec_service.analyze(created["name"])
		return frappe.get_doc("AI Form Template", created["name"])

	def _with_checklist_grid(self, template):
		spec = spec_service.get_spec(template)
		spec["grids"] = [
			{
				"name": "checks",
				"label": "Checks",
				"section": "Checklist",
				"columns": [
					{"name": "item", "label": "Item", "type": "Data", "width": 70},
					{
						"name": "result",
						"label": "Result",
						"type": "Select",
						"options": ["OK", "Not OK"],
						"width": 30,
					},
				],
				"fixed_rows": [{"item": "Formwork clean"}, {"item": "Rebar cover checked"}],
				"blank_rows": 4,
				"allow_add": True,
			}
		]
		spec["pages"][0]["blocks"].append({"type": "grid", "grid": "checks"})
		for field in spec["fields"]:
			if field["name"] == "project_no":
				field["mandatory"] = True
		return spec_forms.save_spec(template.name, json.dumps(spec))

	def test_excel_sheet_to_published_form_without_any_doctype(self):
		before = _custom_doctype_count()
		template = self._excel_form()
		self.assertEqual(
			(template.status, template.storage_mode, template.source_kind),
			("Review Required", "Spec", "Excel"),
		)
		review = self._with_checklist_grid(template)
		self.assertEqual(review["errors"], [])
		# Project configuration sees the spec's fields through the mirrored rows.
		template.reload()
		self.assertIn("checks", [row.fieldname for row in template.fields])
		self.assertIn("CONCRETE POUR CHECKLIST", spec_forms.preview(template.name))
		spec_forms.publish(template.name)
		template.reload()
		self.assertEqual(template.status, "Published")
		self.assertIsNone(template.generated_doctype)
		self.assertEqual(_custom_doctype_count(), before)
		with self.assertRaises(frappe.ValidationError):
			from ai_form_builder.api.generator import generate_doctype

			generate_doctype(template.name, "AFB Test Never")

	def test_record_round_trip_rules_and_pdf(self):
		template = self._excel_form("AFB-T-SPEC2")
		self._with_checklist_grid(template)
		spec_forms.publish(template.name)
		# Not enabled for the project yet.
		with self.assertRaises(frappe.ValidationError):
			records.save({"form_template": template.name, "reference_name": "AFB-A", "values": {}})
		project_configuration.set_form_enabled(PROJECT, "AFB-A", template.name, 1)
		forms = form_registry.get_enabled_forms(PROJECT, "AFB-A", area="Quality")
		spec_form = next(form for form in forms if form["form_template"] == template.name)
		self.assertTrue(spec_form["is_spec"])
		self.assertIsNone(spec_form["doctype"])

		blank = records.get_form(form_template=template.name, reference_name="AFB-A")
		self.assertEqual(
			[row["item"] for row in blank["grids"]["checks"]], ["Formwork clean", "Rebar cover checked"]
		)
		name = records.save(
			{
				"form_template": template.name,
				"reference_name": "AFB-A",
				"values": {"date": "2026-10-09", "inspected_by": "Jane Inspector", "unknown": "dropped"},
				"grids": {
					"checks": [
						{"item": "Tampered text", "result": "OK"},
						{"item": "Rebar cover checked", "result": "Not OK"},
						{"item": "Extra line", "result": "OK"},
					]
				},
			}
		)
		record = frappe.get_doc("AI Form Record", name)
		self.assertTrue(name.startswith("AFB-T-SPEC2-"))
		self.assertEqual(record.values_dict(), {"date": "2026-10-09", "inspected_by": "Jane Inspector"})
		rows = record.grid_rows_dict()["checks"]
		# Pre-printed text is re-imposed; added lines are kept.
		self.assertEqual([r["item"] for r in rows], ["Formwork clean", "Rebar cover checked", "Extra line"])
		self.assertTrue(record.configuration_snapshot)
		# A record that answered only the first printed line still shows them all.
		short = records.save(
			{
				"form_template": template.name,
				"reference_name": "AFB-A",
				"grids": {"checks": [{"result": "OK"}]},
			}
		)
		lines = records.get_form(name=short)["grids"]["checks"]
		self.assertEqual(
			[(r["item"], r.get("result")) for r in lines],
			[("Formwork clean", "OK"), ("Rebar cover checked", None)],
		)
		# Completing requires the mandatory Project No.
		with self.assertRaises(frappe.MandatoryError):
			records.save({"status": "Completed", "values": {"date": "2026-10-09"}}, name=name)
		records.save(
			{"status": "Completed", "values": {"project_no": "P-974", "date": "2026-10-09"}}, name=name
		)
		with self.assertRaises(frappe.ValidationError):
			records.save(
				{"values": {"result": "Maybe"}, "grids": {"checks": [{"result": "Maybe"}]}}, name=name
			)

		file_url = records.download_pdf(name)
		content = frappe.get_doc("File", {"file_url": file_url}).get_content()
		with fitz.open(stream=content, filetype="pdf") as pdf:
			text = " ".join(page.get_text() for page in pdf)
		for expected in ("CONCRETE POUR CHECKLIST", "P-974", "Formwork clean", "Not OK"):
			self.assertIn(expected, text)

		# A later revision never changes how this record prints.
		from ai_form_builder.api import forms as forms_api

		revision = forms_api.create_revision(template.name)
		self.assertEqual(frappe.db.get_value("AI Form Template", revision, "status"), "Draft")
		template.reload()
		self.assertEqual(frappe.get_doc("AI Form Record", name).spec(), json.loads(template.spec_json))

	def test_pdf_sheet_goes_to_the_model(self):
		spec = {
			"title": "AFB PDF FORM",
			"fields": [{"name": "who", "label": "Who", "type": "Data"}],
			"pages": [{"blocks": [{"type": "table", "rows": [[{"text": "Who"}, {"field": "who"}]]}]}],
		}
		created = template_api.create_template_from_upload(
			"AFB Test PDF", "AFB-T-PDF", _pdf_url("afb-spec-pdf")
		)

		class Provider:
			def generate_spec(self, kind, path, title=None, draft=None, hints=None):
				self.kind = kind
				return form_spec.normalise(spec)

		provider = Provider()
		config = {
			"provider": "Anthropic",
			"api_key": "sk-ant-test",
			"model": "claude-opus-5-5",
			"source": "test",
		}
		with (
			patch("ai_form_builder.integrations.llm.get_config", return_value=config),
			patch("ai_form_builder.integrations.llm.get_provider", return_value=provider),
		):
			spec_service.analyze(created["name"])
		template = frappe.get_doc("AI Form Template", created["name"])
		self.assertEqual(provider.kind, "PDF")
		self.assertEqual((template.status, template.ai_model), ("Review Required", "claude-opus-5-5"))
		self.assertEqual(spec_service.get_spec(template)["title"], "AFB PDF FORM")

	def test_pdf_without_any_ai_key_fails_clearly(self):
		created = template_api.create_template_from_upload(
			"AFB Test PDF2", "AFB-T-PDF2", _pdf_url("afb-spec-pdf2")
		)
		with (
			patch("ai_form_builder.integrations.llm.get_config", return_value=NO_AI),
			self.assertRaises(frappe.ValidationError),
		):
			spec_service.analyze(created["name"])
		self.assertEqual(
			frappe.db.get_value("AI Form Template", created["name"], "status"), "Analysis Failed"
		)

	def test_export_then_import_moves_a_form_between_sites(self):
		template = self._excel_form("AFB-T-EXP")
		payload = spec_forms.export_spec(template.name)
		payload["template_code"] = "AFB-T-IMP"
		imported = frappe.get_doc("AI Form Template", spec_forms.import_spec(json.dumps(payload)))
		self.assertEqual(imported.status, "Review Required")
		self.assertEqual(spec_service.get_spec(imported), spec_service.get_spec(template))

	def test_renderer_escapes_values_and_rejects_foreign_images(self):
		spec = form_spec.clean(
			{
				"title": "T",
				"fields": [
					{"name": "who", "label": "Who"},
					{"name": "sign", "label": "Sign", "type": "Signature"},
				],
				"pages": [{"blocks": [{"type": "table", "rows": [[{"field": "who"}, {"field": "sign"}]]}]}],
			}
		)
		html = spec_renderer.render_html(spec, {"who": "<script>x</script>", "sign": 'http://evil/"x'}, {})
		self.assertNotIn("<script>x", html)
		self.assertNotIn("evil", html)
